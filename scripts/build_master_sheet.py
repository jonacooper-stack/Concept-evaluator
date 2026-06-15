#!/usr/bin/env python3
"""Build the master comparison sheet (.csv + .xlsx) for the run.

Edit ROWS below, then: python3 scripts/build_master_sheet.py
Emits deliverables/<date>/00-Master-Comparison-Sheet.csv (+ .xlsx if openpyxl).
"""
import csv, os

DATE = "2026-06-15"
OUTDIR = f"deliverables/{DATE}"
HEADER = [
    "Rank", "Concept", "Forcing-function type", "What it is (one line)",
    "MustHave/10", "Competition/10", "Objectives Score/100",
    "CFO", "CMO", "COO", "Legal", "CTO", "Competitive",
    "Lowest sub-score", "Legal risk-gate", "Tier", "Verdict",
]
# Sorted by Objectives Score (desc). All three reached full council.
ROWS = [
    ["1", "Permit & License Renewal Management", "Operational continuity (license to operate)",
     "System of record + done-for-you renewals so a multi-location operator never loses a site to a lapsed permit",
     "9", "6", "79",
     "8.2", "7.4", "7.1", "7.5", "7.5", "6.9",
     "5 (CompAn encroachment)", "SERIOUS-BUT-MANAGEABLE", "B (top of band)", "REFINE -> GO (narrow to one vertical)"],
    ["2", "Certified-Payroll & Prevailing-Wage Compliance", "Critical input to getting paid (+ regulatory mandate)",
     "Weekly certified-payroll filing for public-works subcontractors so their government progress payment is never withheld",
     "9", "5", "76.5",
     "7.6", "7.8", "7.4", "7.7", "7.5", "5.8",
     "4 (CompAn diff/encroach)", "SERIOUS-BUT-MANAGEABLE", "B (top of band)", "REFINE -> GO (pick a slice Miter won't chase)"],
    ["3", "MoCRA Cosmetics Compliance Program", "Regulatory mandate (FDA / MoCRA)",
     "Done-for-you FDA registration/listing/safety/adverse-event program for indie beauty brands",
     "8", "4", "69",
     "7.9", "7.0", "7.6", "7.3", "7.9", "4.7",
     "3 (CompAn diff/encroach/response)", "SERIOUS-BUT-MANAGEABLE", "B (weak, bottom of band)", "REFINE (lean NO-GO as written)"],
]
# Screened-out field (reached triage only) for completeness — top excluded + the two FATALs.
SCREENED = [
    ["—", "RIA Compliance-as-a-Service (1st council alternate)", "Regulatory", "SEC-forced compliance retainer for small RIAs", "", "", "triage 7.5", "", "", "", "", "", "", "", "MINOR/SERIOUS", "screened (not councilled)", "alternate"],
    ["—", "ISO-17025 Calibration-Certificate Mgmt", "Critical input (production)", "Expired calibration cert halts shipments; cert system of record", "", "", "triage 7.6", "", "", "", "", "", "", "", "NONE", "screened", "candidate"],
    ["—", "Supplier-Quality Doc Management", "Contractual (prime-mandated)", "PPAP/cert docs for Tier-2/3 manufacturers; no docs = no PO", "", "", "triage 7.5", "", "", "", "", "", "", "", "NONE", "screened (one-pager drafted)", "candidate"],
    ["—", "Cyber-Insurance Readiness", "Risk mitigation (insurer-mandated)", "Attested controls SMBs need to bind/renew cyber insurance", "", "", "triage 7.8", "", "", "", "", "", "", "", "MINOR", "screened", "candidate (funded insurtech competition)"],
    ["—", "Retail PIM / Data Syndication", "Critical input", "Product-data syndication to retailer portals", "", "", "FATAL", "", "", "", "", "", "", "", "—", "screened-out", "FATAL: Salsify well-funded strong-tech + heavy integration"],
    ["—", "EDI Integration-as-a-Service", "Critical input", "EDI onboarding to big-box retailers", "", "", "FATAL", "", "", "", "", "", "", "", "—", "screened-out", "FATAL: SPS Commerce dominant strong-tech incumbent"],
]


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    csv_path = os.path.join(OUTDIR, "00-Master-Comparison-Sheet.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["MASTER COMPARISON SHEET — Concept Evaluator run " + DATE])
        w.writerow(["SURFACED TOP 3 (full six-member council + two-directional red-team), sorted by Objectives Score"])
        w.writerow(HEADER)
        for r in ROWS:
            w.writerow(r)
        w.writerow([])
        w.writerow(["SCREENED FIELD (reached triage; not advanced to full council) + notable screen-outs"])
        w.writerow(HEADER)
        for r in SCREENED:
            w.writerow(r)
    print(f"[ok] {csv_path}")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "Master Comparison"
        ws.append(["MASTER COMPARISON SHEET — Concept Evaluator run " + DATE])
        ws.append(["SURFACED TOP 3 (full council + two-directional red-team), sorted by Objectives Score"])
        ws.append(HEADER)
        for r in ROWS: ws.append(r)
        ws.append([])
        ws.append(["SCREENED FIELD (triage only) + notable screen-outs"])
        ws.append(HEADER)
        for r in SCREENED: ws.append(r)
        hdr_fill = PatternFill("solid", fgColor="1F3864")
        for row_idx in (3, 9):  # the two header rows
            for c in ws[row_idx]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
                c.alignment = Alignment(vertical="top", wrap_text=True)
        widths = [6, 38, 30, 50, 10, 12, 16, 6, 6, 6, 7, 6, 11, 22, 22, 20, 38]
        for i, wd in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = wd
        xlsx_path = os.path.join(OUTDIR, "00-Master-Comparison-Sheet.xlsx")
        wb.save(xlsx_path)
        print(f"[ok] {xlsx_path}")
    except Exception as e:
        print(f"[warn] openpyxl not available ({e}); CSV only.")


if __name__ == "__main__":
    main()
