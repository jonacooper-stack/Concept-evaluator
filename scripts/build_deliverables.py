#!/usr/bin/env python3
"""
build_deliverables.py — render founder-facing deliverables (doc 14) as attachments.

Two modes:
  * Markdown prose (Concept Dossier / Scorecard / summary):  .md in  ->  .md, .txt, .docx
    (and optional .pdf with --pdf).
  * Tabular data (master comparison sheet / scoreboard):      .csv in ->  .csv, .xlsx.

The point: produce files the founders can open and forward as attachments — Word
(.docx), text (.txt), Markdown (.md), and Excel (.xlsx). No Google Drive, no GitHub
digging required.

Usage:
    python3 scripts/build_deliverables.py <input.md>  [--outdir deliverables/] [--name NAME] [--pdf]
    python3 scripts/build_deliverables.py <input.csv> [--outdir deliverables/] [--name NAME]

Backends (install once if missing): pip install python-docx openpyxl  (reportlab only if --pdf).
Parses headings, bullet/numbered lists, pipe tables, bold (**...**), rules, code fences.
"""
import argparse
import os
import re
import shutil
import subprocess
import sys


def read_md(path):
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


# ---------- plain text ----------
def to_txt(md):
    out = []
    for line in md.splitlines():
        s = line
        s = re.sub(r"^#{1,6}\s*", "", s)          # heading marks
        s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)     # bold
        s = re.sub(r"`([^`]+)`", r"\1", s)          # inline code
        if re.match(r"^\s*```", s):                  # fence lines
            continue
        out.append(s)
    return "\n".join(out)


# ---------- docx ----------
def to_docx(md, docx_path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except Exception as e:
        print(f"[warn] python-docx not available ({e}); skipping .docx. "
              f"Install with: pip install python-docx", file=sys.stderr)
        return False

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md.splitlines()
    i = 0
    in_code = False
    code_buf = []

    def add_runs(paragraph, text):
        # split on **bold**
        parts = re.split(r"(\*\*.+?\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                r = paragraph.add_run(part[2:-2])
                r.bold = True
            else:
                part = re.sub(r"`([^`]+)`", r"\1", part)
                paragraph.add_run(part)

    while i < len(lines):
        line = lines[i]

        if re.match(r"^\s*```", line):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # table block
        if "|" in line and re.match(r"^\s*\|?.*\|.*$", line) and line.strip().startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            rows = []
            for tl in tbl_lines:
                if re.match(r"^\s*\|?[\s:|-]+\|?\s*$", tl):  # separator row
                    continue
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncol = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncol)
                table.style = "Light Grid Accent 1"
                for r_idx, r in enumerate(rows):
                    cells = table.add_row().cells
                    for c_idx in range(ncol):
                        txt = r[c_idx] if c_idx < len(r) else ""
                        txt = re.sub(r"\*\*(.+?)\*\*", r"\1", txt)
                        cells[c_idx].text = txt
                        if r_idx == 0:
                            for par in cells[c_idx].paragraphs:
                                for run in par.runs:
                                    run.bold = True
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)), level=min(level, 4))
            i += 1
            continue

        if re.match(r"^\s*---+\s*$", line):
            doc.add_paragraph().add_run("_" * 40)
            i += 1
            continue

        mb = re.match(r"^\s*[-*]\s+(.*)$", line)
        if mb:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, mb.group(1))
            i += 1
            continue

        mn = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if mn:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, mn.group(1))
            i += 1
            continue

        if line.strip() == "":
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1

    doc.save(docx_path)
    return True


# ---------- pdf ----------
def _pdf_via_libreoffice(docx_path, pdf_path, outdir):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not (soffice and os.path.exists(docx_path)):
        return False
    env = dict(os.environ)
    env.setdefault("HOME", "/tmp")
    try:
        subprocess.run(
            [soffice, "--headless",
             "-env:UserInstallation=file:///tmp/lo_deliverables_profile",
             "--convert-to", "pdf:writer_pdf_Export", "--outdir", outdir, docx_path],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            timeout=180, env=env,
        )
        produced = os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")
        if os.path.exists(produced):
            if os.path.abspath(produced) != os.path.abspath(pdf_path):
                shutil.move(produced, pdf_path)
            return True
    except Exception as e:
        print(f"[warn] LibreOffice PDF conversion failed ({e}); trying reportlab.",
              file=sys.stderr)
    return False


def _esc(text):
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`([^`]+)`", r'<font face="Courier">\1</font>', text)
    return text


def _pdf_via_reportlab(md, pdf_path):
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                        Table, TableStyle, Preformatted, HRFlowable)
    except Exception as e:
        print(f"[warn] reportlab not available ({e}); skipping .pdf.", file=sys.stderr)
        return False

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10.5, leading=15, spaceAfter=6)
    code = ParagraphStyle("code", parent=styles["Code"], fontSize=8.5, leading=11)
    h = [ParagraphStyle(f"h{n}", parent=styles[f"Heading{min(n,4)}"],
                        spaceBefore=12, spaceAfter=6) for n in range(1, 5)]

    flow = []
    lines = md.splitlines()
    i, in_code, code_buf = 0, False, []
    while i < len(lines):
        line = lines[i]
        if re.match(r"^\s*```", line):
            if in_code:
                flow.append(Preformatted("\n".join(code_buf), code)); code_buf = []
            in_code = not in_code
            i += 1; continue
        if in_code:
            code_buf.append(line); i += 1; continue

        if line.strip().startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tl = lines[i]
                if not re.match(r"^\s*\|?[\s:|-]+\|?\s*$", tl):
                    tbl.append([_esc(c.strip()) for c in tl.strip().strip("|").split("|")])
                i += 1
            if tbl:
                ncol = max(len(r) for r in tbl)
                data = [[Paragraph(c, body) for c in (r + [""] * (ncol - len(r)))] for r in tbl]
                t = Table(data, repeatRows=1, hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3864")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]))
                flow.append(t); flow.append(Spacer(1, 8))
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flow.append(Paragraph(_esc(m.group(2)), h[min(len(m.group(1)), 4) - 1]))
            i += 1; continue
        if re.match(r"^\s*---+\s*$", line):
            flow.append(HRFlowable(width="100%", thickness=0.6, color=colors.grey,
                                   spaceBefore=6, spaceAfter=6))
            i += 1; continue
        mb = re.match(r"^\s*[-*]\s+(.*)$", line)
        if mb:
            flow.append(Paragraph("&bull;&nbsp;" + _esc(mb.group(1)), body)); i += 1; continue
        mn = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if mn:
            flow.append(Paragraph(f"{mn.group(1)}.&nbsp;" + _esc(mn.group(2)), body)); i += 1; continue
        if line.strip() == "":
            flow.append(Spacer(1, 4)); i += 1; continue
        flow.append(Paragraph(_esc(line), body)); i += 1

    SimpleDocTemplate(pdf_path, pagesize=letter,
                      leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                      topMargin=0.9 * inch, bottomMargin=0.9 * inch).build(flow)
    return True


def to_pdf(md, docx_path, pdf_path, outdir):
    if _pdf_via_libreoffice(docx_path, pdf_path, outdir):
        return True
    if _pdf_via_reportlab(md, pdf_path):
        return True
    print("[warn] No PDF backend available (LibreOffice/reportlab); skipping .pdf.",
          file=sys.stderr)
    return False


# ---------- xlsx (master comparison sheet / scoreboard) ----------
def to_xlsx_from_csv(csv_path, xlsx_path):
    try:
        import csv as _csv
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except Exception as e:
        print(f"[warn] openpyxl not available ({e}); skipping .xlsx (.csv still written). "
              f"Install with: pip install openpyxl", file=sys.stderr)
        return False
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(_csv.reader(f))
    if not rows:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Scoreboard"
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if r == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
    ncol = len(rows[0])
    for c in range(1, ncol + 1):
        width = max((len(str(row[c - 1])) for row in rows if c - 1 < len(row)), default=10)
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 48)
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="input Markdown (.md) or CSV (.csv) file")
    ap.add_argument("--outdir", default="deliverables", help="output directory")
    ap.add_argument("--name", default=None, help="base name for outputs")
    ap.add_argument("--pdf", action="store_true", help="also emit a .pdf (off by default)")
    args = ap.parse_args()

    os.makedirs(args.outdir, exist_ok=True)
    base = args.name or os.path.splitext(os.path.basename(args.input))[0]

    # Tabular mode: CSV -> CSV + XLSX (master comparison sheet / scoreboard)
    if args.input.lower().endswith(".csv"):
        csv_out = os.path.join(args.outdir, base + ".csv")
        if os.path.abspath(csv_out) != os.path.abspath(args.input):
            shutil.copyfile(args.input, csv_out)
        print(f"[ok] {csv_out}")
        xlsx_path = os.path.join(args.outdir, base + ".xlsx")
        if to_xlsx_from_csv(args.input, xlsx_path):
            print(f"[ok] {xlsx_path}")
        return

    # Prose mode: MD -> MD + TXT + DOCX (+ optional PDF)
    md = read_md(args.input)
    md_out = os.path.join(args.outdir, base + ".md")
    if os.path.abspath(md_out) != os.path.abspath(args.input):
        shutil.copyfile(args.input, md_out)
    print(f"[ok] {md_out}")

    txt_path = os.path.join(args.outdir, base + ".txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(to_txt(md))
    print(f"[ok] {txt_path}")

    docx_path = os.path.join(args.outdir, base + ".docx")
    if to_docx(md, docx_path):
        print(f"[ok] {docx_path}")

    if args.pdf:
        pdf_path = os.path.join(args.outdir, base + ".pdf")
        if to_pdf(md, docx_path, pdf_path, args.outdir):
            print(f"[ok] {pdf_path}")


if __name__ == "__main__":
    main()
