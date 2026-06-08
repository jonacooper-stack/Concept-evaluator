#!/usr/bin/env python3
"""Assemble the UnionComm council evaluation into a single .docx."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path("/home/user/Concept-evaluator/evaluations/unioncomm")
OUT = BASE / "UnionComm-Council-Evaluation.docx"

SECTIONS = [
    ("00-clean-room-one-pager.md", "Clean-Room Concept Brief (the blind input given to every reviewer)"),
    ("07-objectives-scorecard.md", "Objectives Scorecard (Doc-01 weighted rubric)"),
    ("01-cfo-review.md", "CFO Review"),
    ("02-cmo-review.md", "CMO Review"),
    ("03-coo-review.md", "COO Review"),
    ("04-legal-review.md", "Legal / Regulatory Review"),
    ("05-cto-review.md", "CTO Review"),
    ("06-pm-synthesis.md", "Product Manager Synthesis & Red-Team"),
    ("08-scoreboard.md", "Scoreboard & Verdict"),
]

doc = Document()

# Base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def is_table_sep(line):
    return bool(re.match(r"^\s*\|[\s:|-]+\|\s*$", line))


def add_markdown_table(block):
    rows = [
        [c.strip() for c in row.strip().strip("|").split("|")]
        for row in block if not is_table_sep(row)
    ]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            text = cell.replace("**", "")
            para = table.rows[i].cells[j].paragraphs[0]
            run = para.add_run(text)
            if i == 0:
                run.bold = True
    doc.add_paragraph()


def render_text(text):
    """Render a block of plain/markdown text into the doc."""
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # markdown table
        if is_table_row(line):
            block = []
            while i < len(lines) and is_table_row(lines[i]):
                block.append(lines[i])
                i += 1
            add_markdown_table(block)
            continue

        if not stripped:
            i += 1
            continue

        # markdown headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=2)
        # bullet
        elif stripped.startswith("- "):
            p = doc.add_paragraph(stripped[2:].replace("**", ""), style="List Bullet")
            _ = p
        # numbered list
        elif re.match(r"^\d+\.\s", stripped):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", stripped).replace("**", ""),
                              style="List Number")
        # ALL-CAPS section header line (council doc sections)
        elif (stripped == stripped.upper() and len(stripped) <= 70
              and re.search(r"[A-Z]", stripped) and not stripped.startswith("|")):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        else:
            doc.add_paragraph(stripped.replace("**", ""))
        i += 1


# ---- Title page ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Business-Concept Council Evaluation")
r.bold = True
r.font.size = Pt(22)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Member Engagement & Dues-Collection Overlay for Labor Unions")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Independent five-expert council (CFO · CMO · COO · Legal · CTO) + PM synthesis\n"
             "Evaluation date: 2026-06-06").italic = True

doc.add_paragraph()
verdict = doc.add_paragraph()
verdict.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = verdict.add_run("VERDICT: REFINE — near-miss (does not clear the bar)")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(0x99, 0x33, 0x00)

summ = doc.add_paragraph()
summ.alignment = WD_ALIGN_PARAGRAPH.CENTER
summ.add_run("Council mean 6.48 / 10  ·  Objectives 69 / 100  ·  Lowest sub-score 3 (Legal)").italic = True

doc.add_page_break()

# ---- Sections ----
for fname, heading in SECTIONS:
    path = BASE / fname
    if not path.exists():
        continue
    doc.add_heading(heading, level=1)
    render_text(path.read_text())
    doc.add_page_break()

doc.save(OUT)
print("Wrote", OUT)
