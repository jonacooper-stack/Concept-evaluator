#!/usr/bin/env python3
"""Build a shareable Word report for the 2026-06-09 concept-evaluation run."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# --- base styling ---
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

def h1(t):
    p = doc.add_heading(t, level=1)
    return p

def h2(t):
    return doc.add_heading(t, level=2)

def para(t, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    return p

def bullet(t):
    doc.add_paragraph(t, style="List Bullet")

# --- title ---
title = doc.add_heading("Concept Evaluation — Final Report", level=0)
sub = doc.add_paragraph()
sub.add_run("Run date: 2026-06-09   ·   Branch: claude/stoic-thompson-48nvvx").italic = True

# --- 1. bar ---
h1("1. What was asked, and the bar")
para("Run the council evaluator until at least TWO concepts each score, on honest "
     "red-teamed numbers:")
bullet("Objectives weighted total (doc-01) > 75 / 100, AND")
bullet("every one of the five expert averages (CFO, CMO, COO, Legal, CTO) >= 7.5 / 10.")
para("Operator constraints honored this run:")
bullet("Reference the prior corpus so nothing is re-evaluated — a ledger of ~45 concepts "
       "across 10 prior branches was built first; all 14 new candidates were checked clear of it.")
bullet("Cap of 5 total council passes (cost control).")
bullet("Surface the top concepts even if none clear the bar.")

# --- 2. outcome ---
h1("2. Headline outcome")
para("5 council passes run. ZERO concepts cleared the bar.", bold=True)
para("Every concept was scored on absolute merit, the mean computed only after all "
     "sub-scores were fixed, and the mandatory PM red-team only ever LOWERED scores. "
     "Per the operating doctrine, a rigorous run that finds zero qualifiers is a SUCCESS, "
     "not a failure — the failure mode would be manufacturing a qualifier at the line. "
     "Nothing was nudged over.")
para("The result matches the prior corpus: across ~45 earlier concepts in 10 branches, "
     "none ever cleared even the stricter native bar; honest means clustered 6.3–7.4. "
     "This run's five honest means clustered 6.58–7.14. The field is consistent.")

# --- 3. top 10 ---
h1("3. Top 10 concepts evaluated this run (ranked)")
para("Council-validated concepts first (red-teamed), then the strongest triage-only "
     "candidates not councilled due to the 5-pass cap.")
rows = [
    ("Rank","Concept","Obj /100","Mean /10","Lowest expert","Stage","Verdict"),
    ("1","ChemSDS","75.0","7.10","Legal 6.5","Full council","FAIL — near-miss (liability tail)"),
    ("2","AccessGuard","72.0","7.14","COO 6.7","Full council","FAIL — near-miss (labor-scaling)"),
    ("3","FranAudit","71.5","6.60","Legal 4.8","Full council","FAIL — UPL fatal-class flaw"),
    ("4","GovRenew","70.5","7.02","Legal 6.0","Full council","FAIL — FCA/UPL + free-DIY WTP"),
    ("5","ProxyDocket","69.5","6.58","CTO 6.1","Full council","FAIL — latent demand; ingestion"),
    ("6","TariffLedger","~74*","—","—","Triage only","not councilled (pass cap)"),
    ("7","EmissionsLedger","~72*","—","—","Triage only","not councilled (pass cap)"),
    ("8","FundFiler","~71*","—","—","Triage only","not councilled (pass cap)"),
    ("9","RenewRail","~70*","—","—","Triage only","not councilled (pass cap)"),
    ("10","AdReg","~68*","—","—","Triage only","not councilled (pass cap)"),
]
t = doc.add_table(rows=len(rows), cols=7)
t.style = "Light Grid Accent 1"
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = t.cell(i, j)
        cell.text = val
        if i == 0:
            for r in cell.paragraphs[0].runs:
                r.bold = True
para("* Triage gut-check estimate only, NOT council-validated. This run's triage→council "
     "gap averaged ~8 points of optimism (AccessGuard triaged 82, councilled 72), so the "
     "true council numbers for #6–#10 would very likely land high-60s/low-70s.", italic=True)
para("Best two by overall strength (objectives + healthiest distribution, setting aside "
     "FranAudit's fatal Legal flaw): ChemSDS and AccessGuard — honest near-misses with the "
     "clearest path to a future GO via cheap validation, not structural redesign.", bold=True)

# --- 4. differentiation ---
h1("4. How the top two differ (customer · market · revenue mechanism)")
para("ChemSDS — customer: small chemical formulators/blenders/manufacturers; market: "
     "OSHA HazCom 2024 / GHS chemical-safety document compliance; revenue: SDS-authoring "
     "+ label + hosted-portal subscription (~$1,800/mo).")
para("AccessGuard — customer: mid-market companies with high-traffic transactional "
     "websites (entry: just-sued); market: ADA Title III web-accessibility litigation "
     "defense; revenue: monitoring + remediation + legal-grade attestation retainer "
     "(~$2,500/mo).")
para("Materially distinct on all three axes.")

# --- 5. evidence spot-check ---
h1("5. Evidence spot-check on the top two")
para("Restating >=8 sub-scores with the inline evidence that earned them (a number, a "
     "named competitor, or shown math):")
h2("ChemSDS")
bullet("CFO Startup-capital-efficiency 9 — launchable for low five figures, no "
       "inventory/equipment, vs enterprise EHS platforms (VelocityEHS, Sovos) built over years/millions.")
bullet("COO Logistics-tractability 9 — zero physical logistics; 100% remote document/data work to any US state.")
bullet("CMO Competitive-air 8 — VelocityEHS/Verisk 3E/Sphera can't profitably serve a "
       "$500–$5K/mo customer; per-hour consultancies won't productize.")
bullet("CMO Customer/demand-evidence 8 — HazCom perennially among OSHA's most-cited "
       "standards; May 2024 final rule, Jan-2026/Jul-2027 staggered deadlines.")
h2("AccessGuard")
bullet("CFO Working-capital 9 — annual contract billed upfront; customer pays a full year before most COGS.")
bullet("CMO Wedge 9 — buyer list is public record: PACER Title III dockets, ~4,000+ "
       "federal filings/yr (Seyfarth/UsableNet), enumerable named defendants.")
bullet("CMO Customer/demand-evidence 9 — accessiBe paid a $1M FTC settlement (2025); "
       "$5K–$25K settlement ranges; multiple-thousand annual filings in NY/FL/CA.")
bullet("Legal Licensing 9 — no occupational license exists for web-accessibility "
       "services in any US state; WCAG is a W3C standard, not a licensed practice.")
para("All restated >=8 scores carry real evidence; none is a vague adjective. FranAudit "
     "and GovRenew are dragged down precisely where their evidence is weakest (Legal "
     "liability), which is why they FAIL rather than pass.")

# --- 6. failure modes ---
h1("6. The five structural walls (this run AND the prior 45 concepts)")
bullet("Software-vs-BPO margin hinge — the defensible work is human and scales linearly.")
bullet("Liability coupling — the value prop and the largest liability are the same thing.")
bullet("Latent vs forced demand — a 'should-comply' deferred until enforcement is acute.")
bullet("Funded / strong-tech incumbent in the lane (constraint #5).")
bullet("UPL / 'this is legal work' — when the deliverable is a legal document or opinion.")

# --- 7. recommendation ---
h1("7. Honest recommendation")
para("A successful, rigorous, zero-qualifier run. Two legitimate next moves (operator's choice):")
bullet("Take the top two near-misses (ChemSDS, AccessGuard) into cheap pre-build "
       "validation — each has a days-to-weeks experiment set in its PM synthesis; if the "
       "numbers land, re-score from real data and either could legitimately cross the bar.")
bullet("OR keep generating in new territory (the held triage candidates TariffLedger / "
       "EmissionsLedger were never councilled) — but only as an explicit operator decision, "
       "never by inflating a packet.")

# --- appendix: full council scoreboard ---
doc.add_page_break()
h1("Appendix — Full-council scoreboard (post red-team)")
srows = [
    ("#","Concept","CFO","CMO","COO","Legal","CTO","Mean","Obj","PASS?"),
    ("1","AccessGuard","7.4","7.5","6.7","7.2","6.9","7.14","72.0","FAIL"),
    ("2","ChemSDS","7.5","7.3","7.6","6.5","6.6","7.10","75.0","FAIL"),
    ("3","FranAudit","7.3","7.7","6.0","4.8","7.2","6.60","71.5","FAIL"),
    ("4","ProxyDocket","7.0","7.2","6.4","6.2","6.1","6.58","69.5","FAIL"),
    ("5","GovRenew","7.5","7.3","6.9","6.0","7.4","7.02","70.5","FAIL"),
]
st = doc.add_table(rows=len(srows), cols=10)
st.style = "Light Grid Accent 1"
for i, row in enumerate(srows):
    for j, val in enumerate(row):
        cell = st.cell(i, j)
        cell.text = val
        if i == 0:
            for r in cell.paragraphs[0].runs:
                r.bold = True
para("")
para("Bar: NONE cleared obj > 75 (ChemSDS touched 75.0 exactly, not >75) and NONE had "
     "all five expert averages >= 7.5. Each concept was held below the bar by at least one "
     "structural drag; the red-team widened, never closed, those gaps. Full uncompressed "
     "packets (clean-room one-pager + five expert reviews + PM synthesis with red-team) for "
     "every concept live in the repo under runs/2026-06-09-stoic/<Concept>/.", italic=True)

out = "runs/2026-06-09-stoic/Concept-Evaluation-Report.docx"
doc.save(out)
print("Saved", out)
